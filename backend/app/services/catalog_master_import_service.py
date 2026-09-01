"""Stateless preview/apply pipeline for the administrative catalog workbook."""

from __future__ import annotations

import hashlib
import hmac
import io
import json
from datetime import datetime, timedelta, timezone
from typing import Any

import jwt
from jwt.exceptions import InvalidTokenError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.config import settings
from app.models import (
    Brand,
    BrandCountryCell,
    CatalogImportBatch,
    Filament,
    FilamentCountryCell,
    FilamentLine,
    Preset,
    PresetPrinter,
    Printer,
)
from app.models.preset import PresetModerationStatus
from app.schemas.catalog_import import CatalogImportDraft, CatalogImportPlanRow
from app.schemas.country_cell import BrandCountryCellBase, FilamentCountryCellBase
from app.schemas.filament import FilamentBase
from app.schemas.preset import PresetBase
from app.services.brand_slug_service import canonicalize_brand_slug, suggest_brand_slug
from app.services.catalog_color_groups import classify_color_group
from app.services.catalog_url_service import choose_filament_slug

MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5_000
CONFIRMATION_MINUTES = 15
CONFIRMATION_TYPE = "admin_catalog_master_import"

SHEET_COLUMNS: dict[str, list[str]] = {
    "Brands": [
        "enabled",
        "brand_key",
        "mode",
        "existing_brand",
        "overwrite",
        "name",
        "slug",
        "description",
        "website",
        "logo_url",
        "logo_bg",
        "currency",
        "price_hidden",
        "verified",
        "active",
        "social_media_urls",
        "shop_links",
    ],
    "Filaments": [
        "enabled",
        "filament_key",
        "brand_key",
        "mode",
        "existing_filament",
        "overwrite",
        "name",
        "line",
        "material_type",
        "color_name",
        "color_hex",
        "ral_code",
        "color_group",
        "color_group_source",
        "diameter",
        "density",
        "drying_required",
        "drying_temperature_c",
        "drying_duration_hours",
        "enclosure_requirement",
        "chamber_temperature_c",
        "bed_adhesives",
        "recommended_nozzle_temp_min",
        "recommended_nozzle_temp_max",
        "recommended_bed_temp_min",
        "recommended_bed_temp_max",
        "required_nozzle_hrc",
        "price_per_kg",
        "spool_weight",
        "empty_spool_weight_g",
        "price_display_unit",
        "availability",
        "description",
        "active",
        "color_type",
        "colors",
        "finish",
        "effects",
        "transparency",
        "additives",
        "property_claims",
        "post_processing_chemicals",
    ],
    "BrandMarkets": [
        "enabled",
        "brand_key",
        "country",
        "overwrite",
        "website",
        "description",
        "social_media_urls",
        "shop_links",
        "currency",
        "published",
    ],
    "FilamentMarkets": [
        "enabled",
        "filament_key",
        "country",
        "overwrite",
        "availability",
        "price",
        "currency",
        "price_display_unit",
        "product_url",
        "purchase_links",
        "market_note",
        "market_color_name",
        "published",
    ],
    "Presets": [
        "enabled",
        "preset_key",
        "filament_key",
        "mode",
        "existing_preset",
        "overwrite",
        "name",
        "description",
        "extruder_temp",
        "bed_temp",
        "flow_rate",
        "fan_speed",
        "retraction_length",
        "retraction_speed",
        "printer_ids",
        "orcaslicer_settings",
        "external_id",
        "active",
    ],
}

_DRAFT_FIELD_BY_SHEET = {
    "Brands": "brands",
    "Filaments": "filaments",
    "BrandMarkets": "brand_markets",
    "FilamentMarkets": "filament_markets",
    "Presets": "presets",
}


class CatalogMasterImportError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _canonical(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")


def draft_digest(draft: CatalogImportDraft) -> str:
    return hashlib.sha256(_canonical(draft.model_dump(mode="json"))).hexdigest()


def plan_digest(rows: list[dict[str, Any]]) -> str:
    return hashlib.sha256(_canonical(rows)).hexdigest()


def issue_confirmation(
    *, user_id: int, draft_hash: str, calculated_plan_digest: str
) -> tuple[str, datetime]:
    expires_at = datetime.now(timezone.utc) + timedelta(minutes=CONFIRMATION_MINUTES)
    token = jwt.encode(
        {
            "type": CONFIRMATION_TYPE,
            "user_id": user_id,
            "draft_digest": draft_hash,
            "plan_digest": calculated_plan_digest,
            "exp": expires_at,
        },
        settings.SECRET_KEY,
        algorithm=settings.ALGORITHM,
    )
    return token, expires_at


def verify_confirmation(
    *, token: str, user_id: int, draft_hash: str, calculated_plan_digest: str
) -> None:
    try:
        payload = jwt.decode(
            token,
            settings.SECRET_KEY,
            algorithms=[settings.ALGORITHM],
            leeway=30,
        )
        if payload.get("type") != CONFIRMATION_TYPE or payload.get("user_id") != user_id:
            raise ValueError("wrong confirmation owner")
        for key, expected in (
            ("draft_digest", draft_hash),
            ("plan_digest", calculated_plan_digest),
        ):
            actual = payload.get(key)
            if not isinstance(actual, str) or not hmac.compare_digest(actual, expected):
                raise ValueError(f"{key} changed")
    except (InvalidTokenError, TypeError, ValueError) as exc:
        raise CatalogMasterImportError(
            "ERR_CATALOG_IMPORT_CONFIRMATION_INVALID",
            "The preview is stale or belongs to another import",
        ) from exc


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def parse_workbook(payload: bytes, filename: str) -> CatalogImportDraft:
    if len(payload) > MAX_WORKBOOK_BYTES:
        raise CatalogMasterImportError("ERR_CATALOG_IMPORT_FILE_TOO_LARGE", "Workbook is too large")
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl exposes several format-specific exceptions
        raise CatalogMasterImportError(
            "ERR_CATALOG_IMPORT_INVALID_XLSX", "Invalid XLSX workbook"
        ) from exc

    parsed: dict[str, list[dict[str, Any]]] = {
        field: [] for field in _DRAFT_FIELD_BY_SHEET.values()
    }
    total_rows = 0
    for sheet_name, draft_field in _DRAFT_FIELD_BY_SHEET.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        headers = [str(cell.value or "").strip() for cell in header_cells]
        if not any(headers):
            continue
        duplicate_headers = {header for header in headers if header and headers.count(header) > 1}
        if duplicate_headers:
            raise CatalogMasterImportError(
                "ERR_CATALOG_IMPORT_INVALID_XLSX",
                f"Duplicate columns in {sheet_name}: {', '.join(sorted(duplicate_headers))}",
            )
        allowed = set(SHEET_COLUMNS[sheet_name])
        unknown = [header for header in headers if header and header not in allowed]
        if unknown:
            raise CatalogMasterImportError(
                "ERR_CATALOG_IMPORT_INVALID_XLSX",
                f"Unknown columns in {sheet_name}: {', '.join(unknown)}",
            )
        for row_number, cells in enumerate(iterator, start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise CatalogMasterImportError(
                    "ERR_CATALOG_IMPORT_FORMULA_NOT_ALLOWED",
                    f"Formula found in {sheet_name} row {row_number}",
                )
            row = {
                header: _cell_value(cell.value)
                for header, cell in zip(headers, cells, strict=False)
                if header
            }
            if not any(value not in (None, "") for value in row.values()):
                continue
            row["_row"] = row_number
            parsed[draft_field].append(row)
            total_rows += 1
            if total_rows > MAX_ROWS:
                raise CatalogMasterImportError(
                    "ERR_CATALOG_IMPORT_TOO_MANY_ROWS", "Workbook has too many rows"
                )
    workbook.close()
    if not parsed["brands"] and not parsed["filaments"]:
        raise CatalogMasterImportError(
            "ERR_CATALOG_IMPORT_EMPTY", "Brands or Filaments must contain at least one row"
        )
    return CatalogImportDraft(filename=filename, **parsed)


def build_template() -> bytes:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Readme"
    readme.append(["FilamentHub catalog master-import"])
    readme.append(["Fill Brands and Filaments. Markets and Presets are optional."])
    readme.append(["mode: auto/create/update; enabled and overwrite: true/false."])
    readme.append(["Lists may be JSON arrays or values separated with semicolons."])
    readme.append(["Nothing is written before the final Apply confirmation."])
    for name, columns in SHEET_COLUMNS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(columns)).coordinate}"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F46E5")
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(
                36, max(12, len(column) + 2)
            )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _present(value: Any) -> bool:
    return value not in (None, "")


def _text(value: Any) -> str | None:
    if not _present(value):
        return None
    text = str(value).strip()
    return text or None


def _bool(value: Any, *, default: bool | None = None) -> bool | None:
    if not _present(value):
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"1", "true", "yes", "y", "да"}:
        return True
    if normalized in {"0", "false", "no", "n", "нет"}:
        return False
    raise ValueError(f"Invalid boolean: {value}")


def _number(value: Any, *, integer: bool = False) -> float | int | None:
    if not _present(value):
        return None
    number = float(str(value).strip().replace(",", "."))
    if integer:
        if not number.is_integer():
            raise ValueError(f"Expected integer: {value}")
        return int(number)
    return number


def _json(value: Any, expected: type) -> Any:
    if not _present(value):
        return None
    if isinstance(value, expected):
        return value
    parsed = json.loads(str(value))
    if not isinstance(parsed, expected):
        raise ValueError(f"Expected {expected.__name__}")
    return parsed


def _list(value: Any) -> list[Any] | None:
    if not _present(value):
        return None
    if isinstance(value, list):
        return value
    text = str(value).strip()
    if text.startswith("["):
        return _json(text, list)
    return [item.strip() for item in text.split(";") if item.strip()]


def _row_number(row: dict[str, Any], fallback: int) -> int:
    try:
        return int(row.get("_row") or fallback)
    except (TypeError, ValueError):
        return fallback


def _enabled(row: dict[str, Any]) -> bool:
    return bool(_bool(row.get("enabled"), default=True))


def _overwrite(row: dict[str, Any]) -> bool:
    return bool(_bool(row.get("overwrite"), default=False))


def _mode(row: dict[str, Any]) -> str:
    value = (_text(row.get("mode")) or "auto").casefold()
    if value not in {"auto", "create", "update"}:
        raise ValueError("mode must be auto, create or update")
    return value


def _jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if hasattr(value, "value"):
        return value.value
    return value


def _changes(
    target: Any,
    payload: dict[str, Any],
    *,
    overwrite: bool,
    comparable: dict[str, Any] | None = None,
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    changes: dict[str, dict[str, Any]] = {}
    conflicts: list[str] = []
    comparable = comparable or {}
    for field, incoming in payload.items():
        current = comparable.get(field, getattr(target, field, None))
        current_json = _jsonable(current)
        incoming_json = _jsonable(incoming)
        if current_json == incoming_json:
            continue
        current_is_empty = current in (None, "", [], {})
        if not overwrite and not current_is_empty:
            conflicts.append(field)
            continue
        changes[field] = {"before": current_json, "after": incoming_json}
    return changes, conflicts


def _error_message(exc: Exception) -> str:
    if isinstance(exc, ValidationError):
        return "; ".join(
            f"{'.'.join(str(part) for part in item['loc'])}: {item['msg']}"
            for item in exc.errors()[:5]
        )
    return str(exc)


def _public_rows(plan: list[dict[str, Any]]) -> list[CatalogImportPlanRow]:
    return [
        CatalogImportPlanRow(
            sheet=item["sheet"],
            row=item["row"],
            key=item.get("key"),
            status=item["status"],
            message=item.get("message"),
            changes=item.get("changes", {}),
        )
        for item in plan
    ]


def summarize(plan: list[dict[str, Any]]) -> dict[str, int]:
    summary = {"create": 0, "update": 0, "noop": 0, "skipped": 0, "error": 0}
    for item in plan:
        summary[item["status"]] += 1
    return summary


async def build_plan(
    db: AsyncSession, draft: CatalogImportDraft, *, admin_user_id: int
) -> list[dict[str, Any]]:
    plan: list[dict[str, Any]] = []
    brands_by_key: dict[str, dict[str, Any]] = {}
    filaments_by_key: dict[str, dict[str, Any]] = {}

    brands = (await db.execute(select(Brand))).scalars().all()
    brands_by_id = {brand.id: brand for brand in brands}
    brands_by_slug = {brand.slug.casefold(): brand for brand in brands}
    brands_by_name = {brand.name.strip().casefold(): brand for brand in brands}
    reserved_brand_names = set(brands_by_name)
    reserved_brand_slugs = set(brands_by_slug)

    for index, row in enumerate(draft.brands, start=2):
        row_number = _row_number(row, index)
        key = _text(row.get("brand_key"))
        item: dict[str, Any] = {"sheet": "Brands", "row": row_number, "key": key}
        try:
            if not _enabled(row):
                item.update(status="skipped", message="Disabled by administrator")
                plan.append(item)
                continue
            if not key:
                raise ValueError("brand_key is required")
            if key in brands_by_key:
                raise ValueError("brand_key must be unique")
            mode = _mode(row)
            existing_ref = _text(row.get("existing_brand"))
            target = None
            if existing_ref:
                target = (
                    brands_by_id.get(int(existing_ref))
                    if existing_ref.isdecimal()
                    else brands_by_slug.get(existing_ref.casefold())
                    or brands_by_name.get(existing_ref.casefold())
                )
            requested_slug = _text(row.get("slug"))
            name = _text(row.get("name"))
            if target is None and requested_slug:
                target = brands_by_slug.get(requested_slug.casefold())
            if target is None and name:
                target = brands_by_name.get(name.casefold())
            if mode == "create" and target is not None:
                raise ValueError("Brand already exists; choose update or auto")
            if mode == "update" and target is None:
                raise ValueError("Existing brand was not found")
            if not name:
                raise ValueError("name is required")

            payload: dict[str, Any] = {"name": name}
            for field in ("description", "website", "logo_url", "logo_bg", "currency"):
                value = _text(row.get(field))
                if value is not None:
                    payload[field] = value
            for field in ("price_hidden", "verified", "active"):
                if _present(row.get(field)):
                    payload[field] = _bool(row.get(field))
            for field in ("social_media_urls", "shop_links"):
                if _present(row.get(field)):
                    payload[field] = (
                        _list(row.get(field))
                        if field == "social_media_urls"
                        else _json(row.get(field), list)
                    )

            if target is None:
                if name.casefold() in reserved_brand_names:
                    raise ValueError("Duplicate brand name in workbook")
                slug = canonicalize_brand_slug(requested_slug or "") if requested_slug else None
                if requested_slug and slug is None:
                    raise ValueError("Invalid brand slug")
                if slug is None:
                    slug = await suggest_brand_slug(db, name)
                if slug in reserved_brand_slugs:
                    raise ValueError("Duplicate brand slug in workbook")
                payload["slug"] = slug
                payload.setdefault("currency", "RUB")
                payload.setdefault("price_hidden", False)
                payload.setdefault("verified", False)
                payload.setdefault("active", True)
                item.update(status="create", payload=payload, target_id=None, changes={})
                reserved_brand_names.add(name.casefold())
                reserved_brand_slugs.add(slug)
            else:
                if requested_slug and requested_slug.casefold() != target.slug.casefold():
                    raise ValueError("Changing an existing brand slug is not supported in import")
                changes, conflicts = _changes(target, payload, overwrite=_overwrite(row))
                if conflicts:
                    raise ValueError(
                        "Existing values differ: "
                        + ", ".join(conflicts)
                        + "; enable overwrite or edit the row"
                    )
                item.update(
                    status="update" if changes else "noop",
                    payload={field: change["after"] for field, change in changes.items()},
                    target_id=target.id,
                    target_updated_at=target.updated_at.isoformat() if target.updated_at else None,
                    changes=changes,
                )
            brands_by_key[key] = item
        except (ValueError, TypeError, json.JSONDecodeError) as exc:
            item.update(status="error", message=_error_message(exc), changes={})
            if key:
                brands_by_key[key] = item
        plan.append(item)

    existing_brand_ids = [
        item["target_id"] for item in brands_by_key.values() if item.get("target_id") is not None
    ]
    existing_filaments = []
    if existing_brand_ids:
        existing_filaments = (
            (
                await db.execute(
                    select(Filament)
                    .options(selectinload(Filament.line))
                    .where(Filament.brand_id.in_(existing_brand_ids))
                )
            )
            .scalars()
            .all()
        )
    filaments_by_id = {filament.id: filament for filament in existing_filaments}
    filament_slug_map = {
        (filament.brand_id, filament.slug.casefold()): filament for filament in existing_filaments
    }
    filament_identity_map = {
        (
            filament.brand_id,
            filament.name.strip().casefold(),
            filament.material_type.strip().casefold(),
            (filament.color_name or "").strip().casefold(),
        ): filament
        for filament in existing_filaments
    }
    seen_filament_identity: set[tuple[str, str, str, str]] = set()

    for index, row in enumerate(draft.filaments, start=2):
        row_number = _row_number(row, index)
        key = _text(row.get("filament_key"))
        item = {"sheet": "Filaments", "row": row_number, "key": key}
        try:
            if not _enabled(row):
                item.update(status="skipped", message="Disabled by administrator")
                plan.append(item)
                continue
            brand_key = _text(row.get("brand_key"))
            if not key or not brand_key:
                raise ValueError("filament_key and brand_key are required")
            if key in filaments_by_key:
                raise ValueError("filament_key must be unique")
            parent = brands_by_key.get(brand_key)
            if parent is None or parent["status"] in {"error", "skipped"}:
                raise ValueError("Referenced brand is missing or invalid")
            mode = _mode(row)
            name = _text(row.get("name"))
            material_type = _text(row.get("material_type"))
            color_name = _text(row.get("color_name"))
            if not name or not material_type:
                raise ValueError("name and material_type are required")
            identity = (
                brand_key,
                name.casefold(),
                material_type.casefold(),
                (color_name or "").casefold(),
            )
            if identity in seen_filament_identity:
                raise ValueError("Duplicate filament identity in workbook")
            seen_filament_identity.add(identity)

            brand_id = parent.get("target_id")
            target = None
            existing_ref = _text(row.get("existing_filament"))
            if brand_id is not None and existing_ref:
                target = (
                    filaments_by_id.get(int(existing_ref))
                    if existing_ref.isdecimal()
                    else filament_slug_map.get((brand_id, existing_ref.casefold()))
                )
                if target is not None and target.brand_id != brand_id:
                    target = None
            if brand_id is not None and target is None:
                target = filament_identity_map.get(
                    (
                        brand_id,
                        name.casefold(),
                        material_type.casefold(),
                        (color_name or "").casefold(),
                    )
                )
            if mode == "create" and target is not None:
                raise ValueError("Filament already exists; choose update or auto")
            if mode == "update" and target is None:
                raise ValueError("Existing filament was not found")

            raw: dict[str, Any] = {"name": name, "material_type": material_type}
            text_fields = (
                "color_name",
                "color_hex",
                "color_group",
                "color_group_source",
                "ral_code",
                "enclosure_requirement",
                "price_display_unit",
                "availability",
                "description",
            )
            for field in text_fields:
                value = _text(row.get(field))
                if value is not None:
                    raw[field] = value
            for field in (
                "diameter",
                "density",
                "drying_temperature_c",
                "drying_duration_hours",
                "chamber_temperature_c",
                "price_per_kg",
                "spool_weight",
                "empty_spool_weight_g",
            ):
                if _present(row.get(field)):
                    raw[field] = _number(row.get(field))
            for field in (
                "recommended_nozzle_temp_min",
                "recommended_nozzle_temp_max",
                "recommended_bed_temp_min",
                "recommended_bed_temp_max",
                "required_nozzle_hrc",
            ):
                if _present(row.get(field)):
                    raw[field] = _number(row.get(field), integer=True)
            for field in ("drying_required", "active"):
                if _present(row.get(field)):
                    raw[field] = _bool(row.get(field))
            for field in (
                "bed_adhesives",
                "additives",
                "property_claims",
                "post_processing_chemicals",
            ):
                if _present(row.get(field)):
                    raw[field] = (
                        _list(row.get(field))
                        if field == "bed_adhesives"
                        else _json(row.get(field), list)
                    )
            visual_keys = ("color_type", "colors", "finish", "effects", "transparency")
            if any(_present(row.get(field)) for field in visual_keys):
                visual: dict[str, Any] = {}
                for field in ("color_type", "finish"):
                    if _present(row.get(field)):
                        visual[field] = _text(row.get(field))
                for field in ("colors", "effects"):
                    if _present(row.get(field)):
                        visual[field] = _list(row.get(field))
                if _present(row.get("transparency")):
                    visual["transparency"] = _bool(row.get("transparency"))
                raw["visual_settings"] = visual
            if "color_hex" in raw and "color_group" not in raw:
                raw["color_group"] = classify_color_group(raw["color_hex"])
                raw["color_group_source"] = "auto"

            validated = FilamentBase(**raw).model_dump(mode="json")
            explicit_fields = set(raw)
            payload = (
                validated
                if target is None
                else {field: validated[field] for field in explicit_fields}
            )
            if _present(row.get("active")):
                payload["active"] = _bool(row.get("active"))
            line_name = _text(row.get("line"))
            comparable = {"line": target.line.name if target and target.line else None}
            if line_name is not None:
                payload["line"] = line_name
            if target is None:
                item.update(
                    status="create",
                    payload=payload,
                    target_id=None,
                    parent_key=brand_key,
                    changes={},
                )
            else:
                changes, conflicts = _changes(
                    target, payload, overwrite=_overwrite(row), comparable=comparable
                )
                if conflicts:
                    raise ValueError(
                        "Existing values differ: "
                        + ", ".join(conflicts)
                        + "; enable overwrite or edit the row"
                    )
                item.update(
                    status="update" if changes else "noop",
                    payload={field: change["after"] for field, change in changes.items()},
                    target_id=target.id,
                    target_updated_at=target.updated_at.isoformat() if target.updated_at else None,
                    parent_key=brand_key,
                    changes=changes,
                )
            filaments_by_key[key] = item
        except (ValueError, TypeError, json.JSONDecodeError, ValidationError) as exc:
            item.update(status="error", message=_error_message(exc), changes={})
            if key:
                filaments_by_key[key] = item
        plan.append(item)

    await _plan_markets(db, draft, plan, brands_by_key, filaments_by_key)
    await _plan_presets(db, draft, plan, filaments_by_key, admin_user_id)
    return plan


async def _plan_markets(
    db: AsyncSession,
    draft: CatalogImportDraft,
    plan: list[dict[str, Any]],
    brands_by_key: dict[str, dict[str, Any]],
    filaments_by_key: dict[str, dict[str, Any]],
) -> None:
    definitions = (
        (
            "BrandMarkets",
            draft.brand_markets,
            "brand_key",
            brands_by_key,
            BrandCountryCell,
            "brand_id",
            BrandCountryCellBase,
        ),
        (
            "FilamentMarkets",
            draft.filament_markets,
            "filament_key",
            filaments_by_key,
            FilamentCountryCell,
            "filament_id",
            FilamentCountryCellBase,
        ),
    )
    for sheet, rows, key_field, parents, model, parent_id_field, schema in definitions:
        seen: set[tuple[str, str]] = set()
        for index, row in enumerate(rows, start=2):
            row_number = _row_number(row, index)
            parent_key = _text(row.get(key_field))
            country = (_text(row.get("country")) or "").upper()
            key = f"{parent_key}:{country}" if parent_key and country else None
            item: dict[str, Any] = {"sheet": sheet, "row": row_number, "key": key}
            try:
                if not _enabled(row):
                    item.update(status="skipped", message="Disabled by administrator")
                    plan.append(item)
                    continue
                if not parent_key or len(country) != 2:
                    raise ValueError(f"{key_field} and a two-letter country are required")
                if (parent_key, country) in seen:
                    raise ValueError("Duplicate market cell in workbook")
                seen.add((parent_key, country))
                parent = parents.get(parent_key)
                if parent is None or parent["status"] in {"error", "skipped"}:
                    raise ValueError("Referenced catalog row is missing or invalid")
                parent_id = parent.get("target_id")
                target = None
                if parent_id is not None:
                    target = await db.scalar(
                        select(model).where(
                            getattr(model, parent_id_field) == parent_id,
                            model.country == country,
                        )
                    )
                raw: dict[str, Any] = {"country": country}
                if sheet == "BrandMarkets":
                    for field in ("website", "description", "currency"):
                        value = _text(row.get(field))
                        if value is not None:
                            raw[field] = value
                    for field in ("social_media_urls", "shop_links"):
                        if _present(row.get(field)):
                            raw[field] = (
                                _list(row.get(field))
                                if field == "social_media_urls"
                                else _json(row.get(field), list)
                            )
                else:
                    for field in (
                        "availability",
                        "currency",
                        "price_display_unit",
                        "product_url",
                        "market_note",
                        "market_color_name",
                    ):
                        value = _text(row.get(field))
                        if value is not None:
                            raw[field] = value
                    if _present(row.get("price")):
                        raw["price"] = _number(row.get("price"))
                    if _present(row.get("purchase_links")):
                        raw["purchase_links"] = _json(row.get("purchase_links"), list)
                if _present(row.get("published")):
                    raw["published"] = _bool(row.get("published"))
                validated = schema(**raw).model_dump(mode="json")
                payload = (
                    validated
                    if target is None
                    else {field: validated[field] for field in raw if field != "country"}
                )
                if target is None:
                    item.update(
                        status="create",
                        payload=payload,
                        target_id=None,
                        parent_key=parent_key,
                        country=country,
                        changes={},
                    )
                else:
                    changes, conflicts = _changes(target, payload, overwrite=_overwrite(row))
                    if conflicts:
                        raise ValueError(
                            "Existing values differ: "
                            + ", ".join(conflicts)
                            + "; enable overwrite or edit the row"
                        )
                    item.update(
                        status="update" if changes else "noop",
                        payload={field: change["after"] for field, change in changes.items()},
                        target_id=target.id,
                        parent_key=parent_key,
                        country=country,
                        target_updated_at=(
                            target.updated_at.isoformat() if target.updated_at else None
                        ),
                        changes=changes,
                    )
            except (ValueError, TypeError, json.JSONDecodeError, ValidationError) as exc:
                item.update(status="error", message=_error_message(exc), changes={})
            plan.append(item)


async def _plan_presets(
    db: AsyncSession,
    draft: CatalogImportDraft,
    plan: list[dict[str, Any]],
    filaments_by_key: dict[str, dict[str, Any]],
    admin_user_id: int,
) -> None:
    all_printer_ids = set((await db.execute(select(Printer.id))).scalars().all())
    seen_keys: set[str] = set()
    for index, row in enumerate(draft.presets, start=2):
        row_number = _row_number(row, index)
        key = _text(row.get("preset_key"))
        item: dict[str, Any] = {"sheet": "Presets", "row": row_number, "key": key}
        try:
            if not _enabled(row):
                item.update(status="skipped", message="Disabled by administrator")
                plan.append(item)
                continue
            filament_key = _text(row.get("filament_key"))
            if not key or not filament_key:
                raise ValueError("preset_key and filament_key are required")
            if key in seen_keys:
                raise ValueError("preset_key must be unique")
            seen_keys.add(key)
            parent = filaments_by_key.get(filament_key)
            if parent is None or parent["status"] in {"error", "skipped"}:
                raise ValueError("Referenced filament is missing or invalid")
            mode = _mode(row)
            target = None
            existing_ref = _text(row.get("existing_preset"))
            external_id = _text(row.get("external_id")) or f"catalog-master:{key}"
            if existing_ref:
                if not existing_ref.isdecimal():
                    raise ValueError("existing_preset must be a numeric ID")
                target = await db.scalar(
                    select(Preset)
                    .options(selectinload(Preset.printer_links))
                    .where(Preset.id == int(existing_ref))
                )
            if target is None and parent.get("target_id") is not None:
                target = await db.scalar(
                    select(Preset)
                    .options(selectinload(Preset.printer_links))
                    .where(
                        Preset.user_id == admin_user_id,
                        Preset.external_id == external_id,
                    )
                )
            if mode == "create" and target is not None:
                raise ValueError("Preset already exists; choose update or auto")
            if mode == "update" and target is None:
                raise ValueError("Existing preset was not found")

            raw: dict[str, Any] = {
                "name": _text(row.get("name")),
                "extruder_temp": _number(row.get("extruder_temp")),
                "bed_temp": _number(row.get("bed_temp")),
                "is_official": False,
                "is_weighted": False,
            }
            for field in ("description",):
                value = _text(row.get(field))
                if value is not None:
                    raw[field] = value
            for field in ("flow_rate", "retraction_length", "retraction_speed"):
                if _present(row.get(field)):
                    raw[field] = _number(row.get(field))
            if _present(row.get("fan_speed")):
                raw["fan_speed"] = _number(row.get("fan_speed"), integer=True)
            if _present(row.get("orcaslicer_settings")):
                raw["orcaslicer_settings"] = _json(row.get("orcaslicer_settings"), dict)
            validated = PresetBase(**raw).model_dump(mode="json")
            allowed = {
                "name",
                "description",
                "extruder_temp",
                "bed_temp",
                "flow_rate",
                "fan_speed",
                "retraction_length",
                "retraction_speed",
                "orcaslicer_settings",
            }
            payload = {field: value for field, value in validated.items() if field in allowed}
            payload["external_id"] = external_id
            if _present(row.get("active")):
                payload["active"] = _bool(row.get("active"))
            printer_ids = [int(value) for value in (_list(row.get("printer_ids")) or [])]
            unknown_printers = sorted(set(printer_ids) - all_printer_ids)
            if unknown_printers:
                raise ValueError("Unknown printer IDs: " + ", ".join(map(str, unknown_printers)))
            if target is None:
                item.update(
                    status="create",
                    payload=payload,
                    target_id=None,
                    parent_key=filament_key,
                    printer_ids=printer_ids,
                    changes={},
                )
            else:
                if (
                    parent.get("target_id") is not None
                    and target.filament_id != parent["target_id"]
                ):
                    payload["filament_id"] = parent["target_id"]
                changes, conflicts = _changes(target, payload, overwrite=_overwrite(row))
                current_printers = sorted(link.printer_id for link in target.printer_links)
                if printer_ids and sorted(printer_ids) != current_printers:
                    if not _overwrite(row) and current_printers:
                        conflicts.append("printer_ids")
                    else:
                        changes["printer_ids"] = {"before": current_printers, "after": printer_ids}
                if conflicts:
                    raise ValueError(
                        "Existing values differ: "
                        + ", ".join(conflicts)
                        + "; enable overwrite or edit the row"
                    )
                item.update(
                    status="update" if changes else "noop",
                    payload={
                        field: change["after"]
                        for field, change in changes.items()
                        if field != "printer_ids"
                    },
                    target_id=target.id,
                    target_updated_at=target.updated_at.isoformat() if target.updated_at else None,
                    parent_key=filament_key,
                    printer_ids=printer_ids if "printer_ids" in changes else None,
                    changes=changes,
                )
        except (ValueError, TypeError, json.JSONDecodeError, ValidationError) as exc:
            item.update(status="error", message=_error_message(exc), changes={})
        plan.append(item)


async def apply_plan(
    db: AsyncSession,
    draft: CatalogImportDraft,
    plan: list[dict[str, Any]],
    *,
    admin_user_id: int,
) -> CatalogImportBatch:
    summary = summarize(plan)
    if summary["error"]:
        raise CatalogMasterImportError(
            "ERR_CATALOG_IMPORT_HAS_ERRORS", "Fix or disable every invalid row before applying"
        )
    batch = CatalogImportBatch(
        filename=draft.filename,
        source_sha256=draft_digest(draft),
        applied_by_user_id=admin_user_id,
        summary=summary,
        plan_snapshot={"rows": [row.model_dump(mode="json") for row in _public_rows(plan)]},
    )
    db.add(batch)
    await db.flush()

    brand_objects: dict[str, Brand] = {}
    for item in (entry for entry in plan if entry["sheet"] == "Brands"):
        if item["status"] in {"error", "skipped"}:
            continue
        if item["status"] == "create":
            brand = Brand(**item["payload"])
            db.add(brand)
            await db.flush()
        else:
            brand = await db.get(Brand, item["target_id"])
            if brand is None:
                raise CatalogMasterImportError("ERR_CATALOG_IMPORT_STALE", "Brand disappeared")
            for field, value in item.get("payload", {}).items():
                setattr(brand, field, value)
        brand_objects[item["key"]] = brand

    filament_objects: dict[str, Filament] = {}
    for item in (entry for entry in plan if entry["sheet"] == "Filaments"):
        if item["status"] in {"error", "skipped"}:
            continue
        brand = brand_objects[item["parent_key"]]
        payload = dict(item.get("payload", {}))
        line_name = payload.pop("line", None)
        # FilamentBase includes the nullable relationship id in its fully
        # materialized create payload. The importer resolves the human-readable
        # line name itself, so never pass the schema default alongside it.
        payload.pop("line_id", None)
        line_id = None
        if line_name:
            line = await db.scalar(
                select(FilamentLine).where(
                    FilamentLine.brand_id == brand.id,
                    func.lower(FilamentLine.name) == line_name.casefold(),
                )
            )
            if line is None:
                line = FilamentLine(brand_id=brand.id, name=line_name)
                db.add(line)
                await db.flush()
            line_id = line.id
        if item["status"] == "create":
            slug = await choose_filament_slug(
                db,
                brand_id=brand.id,
                name=payload["name"],
                color_name=payload.get("color_name"),
                ral_code=payload.get("ral_code"),
                diameter=payload.get("diameter"),
            )
            if slug is None:
                raise CatalogMasterImportError(
                    "ERR_CATALOG_IMPORT_DUPLICATE_FILAMENT", "Filament identity is ambiguous"
                )
            filament = Filament(
                brand_id=brand.id,
                contributed_by_organization_id=None,
                line_id=line_id,
                slug=slug,
                **payload,
            )
            db.add(filament)
            await db.flush()
        else:
            filament = await db.get(Filament, item["target_id"])
            if filament is None:
                raise CatalogMasterImportError("ERR_CATALOG_IMPORT_STALE", "Filament disappeared")
            if "line" in item.get("changes", {}):
                filament.line_id = line_id
            for field, value in payload.items():
                setattr(filament, field, value)
        filament_objects[item["key"]] = filament

    for item in (entry for entry in plan if entry["sheet"] in {"BrandMarkets", "FilamentMarkets"}):
        if item["status"] in {"error", "skipped", "noop"}:
            continue
        is_brand = item["sheet"] == "BrandMarkets"
        parent = (
            brand_objects[item["parent_key"]] if is_brand else filament_objects[item["parent_key"]]
        )
        model = BrandCountryCell if is_brand else FilamentCountryCell
        parent_field = "brand_id" if is_brand else "filament_id"
        if item["status"] == "create":
            market = model(**{parent_field: parent.id}, **item["payload"])
            db.add(market)
        else:
            market = await db.get(model, item["target_id"])
            if market is None:
                raise CatalogMasterImportError(
                    "ERR_CATALOG_IMPORT_STALE", "Market cell disappeared"
                )
            for field, value in item.get("payload", {}).items():
                setattr(market, field, value)

    for item in (entry for entry in plan if entry["sheet"] == "Presets"):
        if item["status"] in {"error", "skipped", "noop"}:
            continue
        filament = filament_objects[item["parent_key"]]
        payload = dict(item.get("payload", {}))
        if item["status"] == "create":
            preset = Preset(
                filament_id=filament.id,
                user_id=admin_user_id,
                created_by_user_id=admin_user_id,
                name=payload.pop("name"),
                source="admin_master_import",
                import_evidence={"type": "admin_master_import", "batch_id": batch.id},
                is_official=False,
                is_weighted=False,
                moderation_status=PresetModerationStatus.APPROVED,
                **payload,
            )
            db.add(preset)
            await db.flush()
        else:
            preset = await db.get(Preset, item["target_id"])
            if preset is None:
                raise CatalogMasterImportError("ERR_CATALOG_IMPORT_STALE", "Preset disappeared")
            for field, value in payload.items():
                setattr(preset, field, value)
            preset.filament_id = filament.id
        if item.get("printer_ids") is not None:
            await db.execute(delete(PresetPrinter).where(PresetPrinter.preset_id == preset.id))
            for index, printer_id in enumerate(item["printer_ids"]):
                db.add(
                    PresetPrinter(
                        preset_id=preset.id,
                        printer_id=printer_id,
                        is_primary=index == 0,
                    )
                )
    await db.flush()
    return batch


__all__ = [
    "CatalogMasterImportError",
    "MAX_WORKBOOK_BYTES",
    "apply_plan",
    "build_plan",
    "build_template",
    "draft_digest",
    "issue_confirmation",
    "parse_workbook",
    "plan_digest",
    "summarize",
    "verify_confirmation",
    "_public_rows",
]
